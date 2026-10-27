#!/usr/bin/env python3
"""Build a Word (.docx) report of both calculations.

    python3 src/build_report.py [OUTPUT.docx]        (default: results/report.docx; needs pandoc)

The report is written as Markdown with LaTeX math (results/report.md) and converted with
pandoc, which produces native Word equations. The page is laid out like a LaTeX article: serif
text (Cambria, which pairs with Cambria Math), numbered sections, every equation and every
substituted calculation on its own line with right-hand equation numbers in the Method section,
rule-only tables, "Figure n:" captions below and "Table n:" captions above, centred page numbers.
A reference docx is generated on the fly; the equation numbers, table rules and page numbers are
added afterwards with python-docx. Every number comes from fuel_rod.py / sg_design.py at build
time; figures are regenerated into results/figures. Takes about two minutes.
"""
import math
import os
import re
import subprocess
import sys

import nbformat

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import fuel_rod as fr          # noqa: E402
import sg_design as sg         # noqa: E402
import plots                   # noqa: E402
import flowchart               # noqa: E402

SUB = os.path.join(os.path.dirname(HERE), "results")
NB_DIR = os.path.join(os.path.dirname(HERE), "notebooks")
FIG = os.path.join(SUB, "figures")
SERIF = "Cambria"
MONO = "Courier New"
M_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
TEXT_WIDTH = 12240 - 2 * 1417            # Letter page, 2.5 cm margins, in twips
EQ_COLS = (700, TEXT_WIDTH - 1400, 700)  # equation table: margin, equation, number


# ----------------------------------------------------------------------------- markdown helpers
class MD:
    def __init__(self):
        self.lines = []
        self.neq = self.nfig = self.ntab = 0

    def p(self, text=""):
        self.lines += [text, ""]

    def cont(self, text):
        """Text that continues the sentence after a display equation (no first-line indent)."""
        self.lines += ['::: {custom-style="Continuation"}', text, ":::", ""]

    def h(self, text, level):
        self.lines += ["#" * level + " " + text, ""]

    def eq(self, latex, number=False):
        """One display equation on its own line. Numbered equations get a right-hand (n)."""
        if number:
            self.neq += 1
            latex = latex + rf" \text{{@@eq{self.neq}@@}}"
        self.lines += [f"$$ {latex} $$", ""]
        return self.neq if number else None

    def bullets(self, items):
        self.lines += [f"- {it}" for it in items] + [""]

    def refs(self, items):
        self.lines += ['::: {custom-style="Bibliography"}']
        for i, it in enumerate(items, 1):
            self.lines += [f"\\[{i}\\] {it}", ""]
        self.lines += [":::", ""]

    def table(self, header, rows, caption, widths=None):
        """Pipe table. widths: relative column widths (pandoc reads them from the dash counts)."""
        self.ntab += 1
        esc = lambda v: str(v).replace("|", "\\|")
        widths = widths or [1] * len(header)
        dashes = ["-" * max(3, int(round(60 * w / sum(widths)))) for w in widths]
        self.lines += ["| " + " | ".join(esc(h) for h in header) + " |",
                       "|" + "|".join(dashes) + "|"]
        self.lines += ["| " + " | ".join(esc(v) for v in r) + " |" for r in rows]
        self.lines += ["", f"Table: Table {self.ntab}: {caption}", ""]
        return self.ntab

    def fig(self, path, caption, width="6in"):
        self.nfig += 1
        rel = os.path.relpath(path, SUB)
        self.lines += [f"![Figure {self.nfig}: {caption}]({rel}){{width={width}}}", ""]
        return self.nfig

    def text(self):
        return "\n".join(self.lines)


SUP = str.maketrans("0123456789-", "⁰¹²³⁴⁵⁶⁷⁸⁹⁻")
THIN = " "


def sci(x, digits=2):
    """Prose: 6.9e-4 -> '6.90 × 10⁻⁴'."""
    m, e = f"{x:.{digits}e}".split("e")
    return f"{m} × 10{str(int(e)).translate(SUP)}"


def tsci(x, digits=2):
    """Equations: 6.9e-4 -> '6.90\\times10^{-4}'."""
    m, e = f"{x:.{digits}e}".split("e")
    return rf"{m}\times10^{{{int(e)}}}"


def nl(x, digits=0):
    """Equations: thin-space thousands separators from 10 000 upwards."""
    s = f"{x:,.{digits}f}"
    return s.replace(",", "") if abs(x) < 10000 else s.replace(",", r"\,")


def npr(x, digits=0):
    """Prose: narrow no-break space as thousands separator from 10 000 upwards."""
    s = f"{x:,.{digits}f}"
    return s.replace(",", "") if abs(x) < 10000 else s.replace(",", THIN)


def cite(*nums):
    return "\\[" + ", ".join(str(n) for n in nums) + "\\]"


# ----------------------------------------------------------------------------- reference docx
DEG = r"\,\text{°C}"      # °C inside an equation (a text run renders the same everywhere)


def _plain_fonts(style, name):
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rfonts.attrib.pop(qn(a), None)
    for a in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(a), name)


def reference_docx(path):
    """Pandoc's default reference document restyled like a LaTeX article."""
    default = subprocess.run(["pandoc", "--print-default-data-file", "reference.docx"], capture_output=True, check=True).stdout
    if len(default) < 1000:
        raise RuntimeError("pandoc did not return its default reference.docx")
    with open(path, "wb") as f:
        f.write(default)
    d = Document(path)
    black = RGBColor(0, 0, 0)
    J, C, L = WD_ALIGN_PARAGRAPH.JUSTIFY, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT
    names = {s.name for s in d.styles}

    def st(name):
        return d.styles[name] if name in names else None

    def font(s, size, bold=False, italic=False, name=SERIF):
        if s is None:
            return
        _plain_fonts(s, name)
        s.font.size = Pt(size)
        s.font.bold = bold
        s.font.italic = italic
        s.font.color.rgb = black

    def para(s, align=None, before=None, after=None, indent=None, left=None, keep=None, tab=None):
        if s is None:
            return
        pf = s.paragraph_format
        if align is not None:
            pf.alignment = align
        if before is not None:
            pf.space_before = Pt(before)
        if after is not None:
            pf.space_after = Pt(after)
        if indent is not None:
            pf.first_line_indent = Cm(indent)
        if left is not None:
            pf.left_indent = Cm(left)
        if keep is not None:
            pf.keep_with_next = keep
        if tab is not None:
            pf.tab_stops.add_tab_stop(Cm(tab))

    font(st("Normal"), 11)
    para(st("Normal"), align=L, before=0, after=0, indent=0)
    font(st("Body Text"), 11)
    para(st("Body Text"), align=J, before=0, after=0, indent=0.6)          # LaTeX paragraph indent
    font(st("First Paragraph"), 11)
    para(st("First Paragraph"), align=J, before=6, after=0, indent=0)      # after a heading, table or figure
    cont = d.styles.add_style("Continuation", WD_STYLE_TYPE.PARAGRAPH)
    cont.base_style = d.styles["Body Text"]
    para(cont, align=J, before=0, indent=0)
    font(st("Compact"), 11)                                                # list items and table cells
    para(st("Compact"), align=J, before=0, after=2, indent=0)
    bib = st("Bibliography") or d.styles.add_style("Bibliography", WD_STYLE_TYPE.PARAGRAPH)
    font(bib, 10.5)
    para(bib, align=L, before=0, after=3, indent=-0.9, left=0.9)
    for name, size, before, after, tab in (("Heading 1", 14, 18, 4, 1.0), ("Heading 2", 12, 12, 2, 1.2), ("Heading 3", 11, 10, 0, 1.4)):
        font(st(name), size, bold=True)
        para(st(name), align=L, before=before, after=after, indent=0, keep=True, tab=tab)
    for name in ("Heading 4", "Heading 5", "Heading 6"):
        font(st(name), 11, bold=True, italic=True)
        para(st(name), align=L, before=8, after=0, indent=0, keep=True, tab=1.6)
    font(st("Title"), 17)
    para(st("Title"), align=C, before=18, after=6, indent=0)
    font(st("Subtitle"), 14)
    para(st("Subtitle"), align=C, before=0, after=10, indent=0)
    font(st("Author"), 12)
    para(st("Author"), align=C, before=0, after=4, indent=0)
    font(st("Date"), 12)
    para(st("Date"), align=C, before=0, after=20, indent=0)
    for name in ("Image Caption", "Caption", "Table Caption"):
        font(st(name), 10)
    para(st("Image Caption"), align=C, before=4, after=4, indent=0)
    para(st("Caption"), align=C, before=4, after=4, indent=0)
    para(st("Table Caption"), align=C, before=10, after=4, indent=0)
    para(st("Figure"), align=C, before=8, after=0, indent=0)
    para(st("Captioned Figure"), align=C, before=8, after=0, indent=0)
    for name in ("Source Code", "Verbatim Char"):
        font(st(name), 7.5, name=MONO)
    para(st("Source Code"), align=L, before=2, after=4, indent=0)
    for name in ("Hyperlink", "Footer", "Header"):
        font(st(name), 10)
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2.5)
        s.top_margin = s.bottom_margin = Cm(2.5)
    d.save(path)
    return path


# ----------------------------------------------------------------------------- docx post-processing
TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc",
               "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar", "textDirection",
              "tcFitText", "vAlign", "hideMark"]


def _el(tag, **attrs):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn("w:" + k), str(v))
    return el


def _put(parent, child, order):
    """Insert child where the schema wants it (replacing any element with the same tag)."""
    tag = child.tag.split("}")[1]
    for old in parent.findall(qn("w:" + tag)):
        parent.remove(old)
    later = [qn("w:" + t) for t in order[order.index(tag) + 1:]]
    for i, el in enumerate(list(parent)):
        if el.tag in later:
            parent.insert(i, child)
            return child
    parent.append(child)
    return child


def _borders(tag, spec):
    b = _el(tag)
    for edge, (val, sz) in spec.items():
        e = _el("w:" + edge, val=val)
        if val != "nil":
            e.set(qn("w:sz"), str(sz))
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), "000000")
        b.append(e)
    return b


def _size_math_runs(element, size_pt):
    """Inline equations in table cells take their size from a w:rPr inside each m:r."""
    for r in element.iter("{%s}r" % M_NS):
        rpr = r.find(qn("w:rPr"))
        if rpr is None:
            rpr = OxmlElement("w:rPr")
            mpr = r.find("{%s}rPr" % M_NS)
            r.insert(1 if mpr is not None else 0, rpr)
        for tag in ("w:sz", "w:szCs"):
            for old in rpr.findall(qn(tag)):
                rpr.remove(old)
            rpr.append(_el(tag, val=int(size_pt * 2)))


def booktabs(table):
    """LaTeX-style rules: a line above, a thin line under the header, a line below, nothing vertical."""
    tbl = table._tbl
    tblPr = tbl.tblPr
    _put(tblPr, _el("w:jc", val="center"), TBLPR_ORDER)
    _put(tblPr, _borders("w:tblBorders", {"top": ("single", 8), "left": ("nil", 0), "bottom": ("single", 8),
                                          "right": ("nil", 0), "insideH": ("nil", 0), "insideV": ("nil", 0)}), TBLPR_ORDER)
    size = 7.5 if len(table.columns) >= 10 else 9
    for ri, row in enumerate(table.rows):
        for cell in row.cells:
            if ri == 0:
                _put(cell._tc.get_or_add_tcPr(), _borders("w:tcBorders", {"bottom": ("single", 4)}), TCPR_ORDER)
            for para in cell.paragraphs:
                para.alignment = WD_ALIGN_PARAGRAPH.LEFT
                para.paragraph_format.space_before = Pt(1)
                para.paragraph_format.space_after = Pt(1)
                for run in para.runs:
                    run.font.size = Pt(size)
            _size_math_runs(cell._tc, size)


def _cell(width, content, align):
    tc = _el("w:tc")
    tcPr = _el("w:tcPr")
    tcPr.append(_el("w:tcW", w=width, type="dxa"))
    tcPr.append(_el("w:vAlign", val="center"))
    tc.append(tcPr)
    p = _el("w:p")
    pPr = _el("w:pPr")
    pPr.append(_el("w:spacing", before=80, after=80))
    pPr.append(_el("w:jc", val=align))
    p.append(pPr)
    if content is not None:
        p.append(content)
    tc.append(p)
    return tc


def _equation_table(paras):
    """One borderless 3-column table, one row per display equation, number in the right-hand cell."""
    tbl = _el("w:tbl")
    tblPr = _el("w:tblPr")
    tblPr.append(_el("w:tblW", w=5000, type="pct"))
    tblPr.append(_el("w:jc", val="center"))
    tblPr.append(_borders("w:tblBorders", {e: ("nil", 0) for e in ("top", "left", "bottom", "right", "insideH", "insideV")}))
    tblPr.append(_el("w:tblLayout", type="fixed"))
    mar = _el("w:tblCellMar")
    mar.append(_el("w:left", w=0, type="dxa"))
    mar.append(_el("w:right", w=0, type="dxa"))
    tblPr.append(mar)
    tbl.append(tblPr)
    grid = _el("w:tblGrid")
    for w in EQ_COLS:
        grid.append(_el("w:gridCol", w=w))
    tbl.append(grid)
    for p in paras:
        omp = p.find(".//{%s}oMathPara" % M_NS)
        num = None
        for t in list(omp.iter("{%s}t" % M_NS)):
            m = re.fullmatch(r"\s*@@eq(\d+)@@\s*", t.text or "")
            if m:
                num = m.group(1)
                r = t.getparent()
                r.getparent().remove(r)
        number = None
        if num:
            number = _el("w:r")
            t = _el("w:t")
            t.text = f"({num})"
            number.append(t)
        tr = _el("w:tr")
        tr.append(_cell(EQ_COLS[0], None, "left"))
        tr.append(_cell(EQ_COLS[1], omp, "center"))
        tr.append(_cell(EQ_COLS[2], number, "right"))
        tbl.append(tr)
    return tbl


def _spacer():
    p = _el("w:p")
    pPr = _el("w:pPr")
    pPr.append(_el("w:spacing", before=0, after=0))
    p.append(pPr)
    return p


def equation_tables(doc):
    """Turn every run of consecutive display-equation paragraphs into one equation table."""
    body = doc.element.body
    groups, cur = [], []
    for child in list(body):
        if child.tag == qn("w:p") and child.find(".//{%s}oMathPara" % M_NS) is not None:
            cur.append(child)
        else:
            if cur:
                groups.append(cur)
            cur = []
    if cur:
        groups.append(cur)
    for group in groups:
        tbl = _equation_table(group)
        prev, nxt = group[0].getprevious(), group[-1].getnext()
        if prev is not None and prev.tag == qn("w:tbl"):
            group[0].addprevious(_spacer())
        group[0].addprevious(tbl)
        for p in group:
            body.remove(p)
        if nxt is not None and nxt.tag == qn("w:tbl"):
            tbl.addnext(_spacer())
    return len(groups)


def page_numbers(doc):
    for s in doc.sections:
        s.footer.is_linked_to_previous = False
        p = s.footer.paragraphs[0] if s.footer.paragraphs else s.footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "1"), ("end", None)):
            r = p.add_run()
            r.font.size = Pt(10)
            r.font.name = SERIF
            if kind:
                fc = OxmlElement("w:fldChar")
                fc.set(qn("w:fldCharType"), kind)
                r._r.append(fc)
            elif text.strip() == "PAGE":
                it = OxmlElement("w:instrText")
                it.set(qn("xml:space"), "preserve")
                it.text = text
                r._r.append(it)
            else:
                r.text = text


def finish_docx(path):
    d = Document(path)
    for t in d.tables:
        booktabs(t)
    n = equation_tables(d)
    page_numbers(d)
    d.save(path)
    return n


# ----------------------------------------------------------------------------- notebooks
def notebook_markdown(path, md):
    """The notebook, printed: markdown cells as they are (headings shifted down), code cells as code blocks."""
    nb = nbformat.read(path, as_version=4)
    first = True
    for cell in nb.cells:
        if "plot" in cell.metadata.get("tags", []):
            continue
        src = cell.source.rstrip("\n")
        if cell.cell_type == "markdown":
            lines = src.splitlines()
            if first:                                   # drop the title block, keep the question line
                lines = [l for l in lines if l.startswith("## ")]
                first = False
            out = []
            for l in lines:
                if l.startswith("#"):
                    hashes = len(l) - len(l.lstrip("#"))
                    l = "#" * min(hashes + 1, 4) + re.sub(r"^\s*\d+\.\s+", " ", l[hashes:])   # the report numbers headings
                out.append(l)
            md.lines += out + [""]
        elif cell.cell_type == "code":
            md.lines += ["```python", src, "```", ""]


# ----------------------------------------------------------------------------- 1 Problem
def section_problem(md):
    md.h("Problem", 1)
    md.p("Question 1 asks for the largest linear heat rate $q'$ a PWR fuel rod can run at while the pellet centre stays "
         "below 2000 °C and the clad below 1200 °C. The rod has a 10 mm outer diameter, a 9 mm UO₂ pellet and a 0.5 mm "
         "Zircaloy clad, on a 12.5 mm triangular pitch; it is 4 m long with 3 m active. The coolant is water at 310 °C and "
         "180 bar flowing at 5 m/s. This part is analytical.")
    md.p("Question 2 asks us to design one of the four once-through shell-and-tube steam generators of a 1000 MWe plant "
         "with 32 % efficiency. The primary water flows inside the tubes from 325 °C to 295 °C, countercurrent to the "
         "secondary. Feed water enters 20 K subcooled and leaves as steam at 320 °C and 100 bar. Tubes and shell are SS316. "
         "We have to find (i) the number of tubes, (ii) the tube inner diameter, wall thickness and height, (iii) the shell "
         "diameter and thickness, (iv) the pressure drop on each side and (v) the unladen weight, with the height divided "
         "into at least 20 segments, standard correlations and a flow chart of the scheme. This part is done in code; the "
         "notebooks are listed in Section 7.")


# ----------------------------------------------------------------------------- 2 Method
def section_method(md):
    e = {}
    md.h("Method", 1)
    md.h("Fuel rod (Question 1)", 2)
    md.p("We treat the rod as a long cylinder with uniform heat generation in the pellet and steady radial conduction, "
         "and neglect axial conduction. The pellet diameter, 9 mm, is the same as the clad inner diameter, "
         "$10 - 2 \\times 0.5 = 9$ mm, so the rod as given has no gap. We take ideal pellet–clad contact and look at what a "
         "gap would do at the end.")
    md.p("Per metre of rod the heat $q'$ passes through the coolant film and the clad in series. The two resistances "
         "per unit length are")
    e["rfilm"] = md.eq(r"R_{film} = \frac{1}{h\,\pi d_{co}},", True)
    e["rclad"] = md.eq(r"R_{clad} = \frac{\ln(d_{co}/d_{ci})}{2\pi k_c}.", True)
    md.p("Inside the pellet the volumetric heat generation is $q''' = q'/(\\pi r_p^2)$, and the conduction equation with "
         "a temperature-dependent conductivity is")
    e["cond"] = md.eq(r"\frac{1}{r}\,\frac{d}{dr}\left(k_f\,r\,\frac{dT}{dr}\right) + q''' = 0.", True)
    md.cont("Integrating once with $dT/dr = 0$ at the centre, and once more from the surface $T_s$ to the centre $T_c$, "
            "gives the conductivity integral")
    e["int"] = md.eq(r"\int_{T_s}^{T_c} k_f\,dT = \frac{q'}{4\pi}.", True)
    md.cont(f"If $\\overline{{k}}_f$ is the mean conductivity between $T_s$ and $T_c$, Eq. ({e['int']}) is the same as a "
            "resistance $R_{fuel} = 1/(4\\pi \\overline{k}_f)$. Each temperature limit then gives a linear heat rate, and "
            "the smaller of the two is the answer. The hottest point of the clad is its inner surface, so the clad limit "
            "uses $T_{ci}$:")
    e["lim_f"] = md.eq(r"q'_{fuel} = \frac{T_{c,max} - T_b}{R_{film} + R_{clad} + R_{fuel}},", True)
    e["lim_c"] = md.eq(r"q'_{clad} = \frac{T_{ci,max} - T_b}{R_{film} + R_{clad}}.", True)
    md.p(f"In the fuel limit $T_s$ itself depends on $q'$, because $T_s = T_b + q'(R_{{film}} + R_{{clad}})$. So "
         f"Eq. ({e['int']}) has to be solved by iteration: guess $q'$, find $T_s$, evaluate $4\\pi\\int k_f\\,dT$ up to "
         "2000 °C, and repeat until $q'$ stops changing. It converges in a few passes because the outer resistances are "
         "small next to the pellet.")
    md.p("For the film coefficient we take the interior sub-channel of the triangular lattice. Its equivalent diameter is")
    e["de"] = md.eq(r"D_e = d_{co}\left[\frac{2\sqrt3}{\pi}\left(\frac{p}{d_{co}}\right)^2 - 1\right],", True)
    md.cont(f"and we use the Weisman correlation {cite(1, 2)} for turbulent flow parallel to a triangular rod array, valid "
            "for $1.1 \\le p/d \\le 1.5$,")
    e["nu"] = md.eq(r"\mathrm{Nu} = 0.023\,\mathrm{Re}^{0.8}\,\mathrm{Pr}^{0.4}\left[1.130\,\frac{p}{d_{co}} - 0.2609\right],", True)
    e["h"] = md.eq(r"h = \frac{\mathrm{Nu}\,k}{D_e}.", True)
    md.p(f"Water properties come from the IAPWS-IF97 steam tables {cite(3)} at 310 °C and 180 bar. For the clad we use "
         f"the MATPRO fit for Zircaloy {cite(4)}, with $T$ in K, evaluated at the mean clad temperature,")
    e["kc"] = md.eq(r"k_c = 7.51 + 2.09\times10^{-2}\,T - 1.45\times10^{-5}\,T^2 + 7.67\times10^{-9}\,T^3,", True)
    md.cont(f"and for UO₂ at 95 % of theoretical density the fit of Fink {cite(5)}, with $t = T/1000$ and $T$ in K,")
    e["kf"] = md.eq(r"k_f = \frac{100}{7.5408 + 17.692\,t + 3.6142\,t^2} + \frac{6400}{t^{5/2}}\,e^{-16.35/t}.", True)
    md.cont(f"We integrate Eq. ({e['kf']}) numerically. From 0 °C to the melting point (2840 °C) it gives 89.5 W/cm, close "
            "to the 93 W/cm usually quoted, which is a useful check on the fit.")

    md.h("Steam generator (Question 2)", 2)
    md.p("**Energy balance.** We neglect pump heat and losses. The primary is at 180 bar, the core pressure of "
         "Question 1, and all enthalpies are from IAPWS-IF97. The duty of one steam generator and the two flows are")
    e["eb_q"] = md.eq(r"Q_{SG} = \frac{P_e}{\eta\,N_{loops}},", True)
    e["eb_mp"] = md.eq(r"\dot m_p = \frac{Q_{SG}}{h_{p,in} - h_{p,out}},", True)
    e["eb_ms"] = md.eq(r"\dot m_s = \frac{Q_{SG}}{h_{s,out} - h_{fw}}.", True)
    md.p("**Feasibility.** Before sizing anything we check that the two temperature profiles do not cross. If $Q$ is "
         "the heat transferred from the cold end, the local enthalpies and the approach are")
    e["tq_p"] = md.eq(r"h_p(Q) = h_{p,out} + \frac{Q}{\dot m_p},", True)
    e["tq_s"] = md.eq(r"h_s(Q) = h_{fw} + \frac{Q}{\dot m_s},", True)
    e["tq_dT"] = md.eq(r"\Delta T(Q) = T_p(Q) - T_s(Q),", True)
    md.cont("and $\\Delta T$ must stay positive everywhere. In a once-through unit the tight spot is the start of boiling, "
            "where the secondary temperature stops rising. We ask for at least 10 K there.")
    md.p("**Arrangement.** We take a vertical once-through unit of the Babcock & Wilcox type. The primary water flows "
         "down inside straight tubes. Feed water enters the shell at the bottom and flows up along the tubes through "
         "broached support plates; there are no baffles. From the bottom up the shell side is subcooled liquid, then "
         "nucleate boiling, then post-dryout flow once the quality passes $x_{do}$ = 0.9, and finally superheated steam.")
    md.p(f"**Wall thicknesses.** We use the thin-wall formulas of ASME VIII {cite(6)} with SS316 ($S$ = 115 MPa at about "
         "350 °C, $E$ = 1) and a design pressure $P$ of 1.1 times the operating pressure. The tube wall is rounded up to a "
         "standard BWG gauge. The heads and the tubesheets see the primary pressure:")
    e["tube"] = md.eq(r"t_{tube} = \frac{P\,d_o}{2SE + 0.8P},", True)
    e["shell"] = md.eq(r"t_{shell} = \frac{P\,R}{SE - 0.6P},", True)
    e["head"] = md.eq(r"t_{head} = \frac{P\,R}{2SE - 0.2P},", True)
    e["ts"] = md.eq(r"T_{ts} = \frac{D_s}{3}\sqrt{\frac{P}{S}}.", True)
    md.p("**Tube count and shell-side geometry.** The number of tubes follows from the tube-side velocity $V_t$, and "
         "the shell-side flow area, mass flux and equivalent diameter from the triangular pitch $p$:")
    e["N"] = md.eq(r"N = \frac{\dot m_p}{\rho_p V_t A_i},", True)
    e["afree"] = md.eq(r"A_{free} = \frac{\sqrt3}{2}\,p^2 - \frac{\pi}{4}\,d_o^2,", True)
    e["gs"] = md.eq(r"G_s = \frac{\dot m_s}{N A_{free}},", True)
    e["de2"] = md.eq(r"D_e = d_o\left[\frac{2\sqrt3}{\pi}\left(\frac{p}{d_o}\right)^2 - 1\right].", True)
    md.p(f"**Heat transfer.** In single phase we use Dittus–Boelter {cite(7)},")
    e["db"] = md.eq(r"\mathrm{Nu} = 0.023\,\mathrm{Re}^{0.8}\,\mathrm{Pr}^{n},", True)
    md.cont(f"with $n$ = 0.3 on the cooled tube side and 0.4 on the heated shell side, where Re is formed with $D_e$. For "
            f"subcooled and saturated boiling we use the Chen correlation {cite(8)},")
    e["chen"] = md.eq(r"h_{TP} = F\,h_l + S\,h_{FZ},", True)
    e["hl"] = md.eq(r"h_l = 0.023\,\mathrm{Re}_l^{0.8}\,\mathrm{Pr}_l^{0.4}\,\frac{k_l}{D_e}, \qquad \mathrm{Re}_l = \frac{G(1-x)D_e}{\mu_l},", True)
    e["xtt"] = md.eq(r"\frac{1}{X_{tt}} = \left(\frac{x}{1-x}\right)^{0.9}\left(\frac{\rho_l}{\rho_g}\right)^{0.5}\left(\frac{\mu_g}{\mu_l}\right)^{0.1},", True)
    e["F"] = md.eq(r"F = \begin{cases}1 & 1/X_{tt} \le 0.1\\ 2.35\,(1/X_{tt} + 0.213)^{0.736} & 1/X_{tt} > 0.1,\end{cases}", True)
    e["S"] = md.eq(r"S = \frac{1}{1 + 2.53\times10^{-6}\,(\mathrm{Re}_l F^{1.25})^{1.17}},", True)
    e["hfz"] = md.eq(r"h_{FZ} = 0.00122\,\frac{k_l^{0.79}\,c_{p,l}^{0.45}\,\rho_l^{0.49}}{\sigma^{0.5}\,\mu_l^{0.29}\,h_{fg}^{0.24}\,\rho_g^{0.24}}\,"
                     r"\Delta T_{sat}^{0.24}\,\Delta p_{sat}^{0.75}.", True)
    md.cont("In the subcooled region $x = 0$ and $F = 1$; the convective part acts on $(T_w - T_b)$ and the nucleate part, "
            f"once $T_w > T_{{sat}}$, on $(T_w - T_{{sat}})$. After dryout ($x_{{do}} \\le x \\le 1$) we use "
            f"Dougall–Rohsenow {cite(9)} with the vapour at saturation,")
    e["dr"] = md.eq(r"h_{pdo} = 0.023\,\frac{k_g}{D_e}\left[\frac{G D_e}{\mu_g}\left(x + \frac{\rho_g}{\rho_l}(1-x)\right)\right]^{0.8}\mathrm{Pr}_g^{0.4}.", True)
    md.p(f"**Pressure drop.** The friction factor is Filonenko's for smooth tubes {cite(10)} on both sides, and in the "
         "two-phase region we use the homogeneous density:")
    e["f"] = md.eq(r"f = (0.79\ln\mathrm{Re} - 1.64)^{-2},", True)
    e["rhoh"] = md.eq(r"\frac{1}{\rho_H} = \frac{x}{\rho_g} + \frac{1-x}{\rho_l}.", True)
    md.cont("The shell-side pressure drop of a segment has friction, gravity, acceleration and a support-plate loss "
            "($K$ = 0.5 per plate, one plate per metre):")
    e["dp"] = md.eq(r"\Delta p = f\,\frac{\Delta z}{D_e}\,\frac{G_s^2}{2\rho_H} + \rho_H\,g\,\Delta z + "
                    r"G_s^2\,\Delta\!\left(\frac{1}{\rho_H}\right) + K\,\frac{G_s^2}{2\rho_H}.", True)
    md.cont("On the tube side we use the same friction term with $d_i$, plus $K$ = 1.5 for the tube ends. The static "
            "head gained in the downflow is reported separately.")
    md.p("**Local heat flux.** All resistances are referred to the outer tube surface, with a fouling allowance of "
         "$2.5\\times10^{-5}$ m² K/W on each side:")
    e["rin"] = md.eq(r"R_{in} = \frac{d_o}{d_i}\left(\frac{1}{h_i} + R_{f,i}\right) + \frac{d_o\ln(d_o/d_i)}{2k_w} + R_{f,o}.", True)
    md.cont("Since $h_{FZ}$ depends on the wall temperature, $T_w$ is found in each segment by bisection on")
    e["q"] = md.eq(r"q'' = \frac{T_p - T_w}{R_{in}} = q''_{shell}(T_w),", True)
    md.cont("and then $U = q''/(T_p - T_s)$.")
    md.p("**Marching.** At the bottom both streams are known: the primary leaves at 295 °C and the feed water enters "
         "at its inlet temperature. So we start there and march upward in 40 equal segments. In each segment we evaluate "
         "the coefficients at the mid-segment state, transfer")
    e["march"] = md.eq(r"\Delta Q_i = q''_i\,N\pi d_o\,\Delta z,", True)
    md.cont("update both enthalpies,")
    e["hs"] = md.eq(r"h_{s,i+1} = h_{s,i} + \frac{\Delta Q_i}{\dot m_s}, \qquad h_{p,i+1} = h_{p,i} + \frac{\Delta Q_i}{\dot m_p},", True)
    md.cont("and update the shell pressure. The march stops when the secondary reaches the outlet enthalpy, and the "
            "height is the sum of the segments. Since the flows come from the overall balance, the primary reaches "
            "325 °C at the top by itself; we use this as the check on the energy balance. The shell inlet pressure is "
            "adjusted in three outer passes so that the outlet is at the specified pressure. Where a regime boundary "
            "($h_f$, dryout, $h_g$) falls inside a segment, the segment is split at the boundary.")
    md.p("**Weight.** The dry weight is the metal volume of the tubes, the shell, two hemispherical heads, two "
         "tubesheets less the holes and the support plates (25 mm, half solid), times the density of SS316, plus 10 % "
         f"for nozzles, skirt and internals. Figure {md.nfig + 1} is the flow chart of the scheme.")
    md.fig(flowchart.draw(os.path.join(FIG, "flowchart.png")), "Flow chart of the steam-generator calculation.", "5.0in")
    return e


# ----------------------------------------------------------------------------- 3 Question 1
def section_q1(md, e):
    r = fr.solve_fuel_rod()
    rod, film, w = r["rod"], r["film"], r["film"]["props"]
    t = fr.temperatures(r["q_max"], r)
    tc = fr.temperatures(r["q_clad"], r)
    pc = fr.plant_check(r)
    rows = fr.sensitivities(r)
    pd = film["pd"]
    geo = (2 * math.sqrt(3) / math.pi) * pd ** 2 - 1
    wb = 1.130 * pd - 0.2609
    _, _, _, R_clad_c, k_c_c = fr._outer_chain(rod, r["q_clad"], r["R_film"], 0.0)
    hist = r["history"]
    T_cm = 0.5 * (t["T_clad_out"] + t["T_clad_in"])
    h325 = fr.water_props(325.0, 180.0)["h"] / 1e3
    h295 = fr.water_props(295.0, 180.0)["h"] / 1e3
    qf = r["q_fuel"]
    g1 = fr.solve_fuel_rod(h_gap=1e4)["q_max"] / 1e3
    g2 = fr.solve_fuel_rod(h_gap=5e3)["q_max"] / 1e3

    md.h("Question 1: Maximum permissible linear heat rate", 1)
    md.h("Calculation", 2)
    md.p(f"The pitch-to-diameter ratio is $p/d = 12.5/10 = {pd:.2f}$, so Eq. ({e['de']}) gives")
    md.eq(rf"D_e = 10\left[\frac{{2\sqrt3}}{{\pi}}(1.25)^2 - 1\right] = 10 \times {geo:.4f} = {film['D_e'] * 1e3:.3f}\ \mathrm{{mm}}.")
    md.p(f"At 310 °C and 180 bar the steam tables give $\\rho$ = {w['rho']:.1f} kg/m³, $\\mu$ = {w['mu'] * 1e6:.2f} µPa s, "
         f"$k$ = {w['k']:.4f} W/m K and Pr = {w['Pr']:.3f}. The saturation temperature is {w['T_sat']:.1f} °C, so the coolant "
         f"is {w['T_sat'] - 310:.0f} K subcooled. From Eqs. ({e['nu']}) and ({e['h']}),")
    md.eq(rf"\mathrm{{Re}} = \frac{{\rho V D_e}}{{\mu}} = \frac{{{w['rho']:.1f} \times 5 \times {film['D_e'] * 1e3:.3f}\times10^{{-3}}}}"
          rf"{{{w['mu'] * 1e6:.2f}\times10^{{-6}}}} = {tsci(film['Re'])},")
    md.eq(rf"\mathrm{{Nu}} = 0.023\,({tsci(film['Re'])})^{{0.8}}\,({w['Pr']:.3f})^{{0.4}}\,[1.130 \times 1.25 - 0.2609]")
    md.eq(rf"\mathrm{{Nu}} = 0.023 \times {nl(film['Re'] ** 0.8)} \times {w['Pr'] ** 0.4:.3f} \times {wb:.4f} = {film['Nu']:.0f},")
    md.eq(rf"h = \frac{{\mathrm{{Nu}}\,k}}{{D_e}} = \frac{{{film['Nu']:.0f} \times {w['k']:.4f}}}{{{film['D_e'] * 1e3:.3f}\times10^{{-3}}}} "
          rf"= {nl(film['h'])}\ \mathrm{{W/m^2\,K}},")
    md.eq(rf"R_{{film}} = \frac{{1}}{{h\,\pi d_{{co}}}} = \frac{{1}}{{{nl(film['h'])} \times \pi \times 0.010}} = {tsci(r['R_film'])}\ \mathrm{{K\,m/W}}.")
    md.p(f"The clad is thin and its conductivity changes slowly, so we evaluate Eq. ({e['kc']}) at the mean clad "
         f"temperature of the final solution, {T_cm:.0f} °C ({T_cm + 273.15:.0f} K). That gives $k_c$ = {r['k_clad']:.2f} W/m K "
         f"and, from Eq. ({e['rclad']}),")
    md.eq(rf"R_{{clad}} = \frac{{\ln(10/9)}}{{2\pi \times {r['k_clad']:.2f}}} = \frac{{{math.log(10 / 9):.4f}}}{{{2 * math.pi * r['k_clad']:.1f}}} "
          rf"= {tsci(r['R_clad'])}\ \mathrm{{K\,m/W}}.")
    md.p(f"**Pellet-centre limit.** We start the iteration at 50 kW/m. Each pass computes $T_s$ from the film and clad "
         f"drops and then a new $q'$ from Eq. ({e['int']}) with $T_c$ = 2000 °C (Table {md.ntab + 1}). It settles at "
         f"$q'_{{fuel}}$ = {qf / 1e3:.2f} kW/m with $T_s$ = {t['T_fuel_surface']:.1f} °C. At that point")
    md.eq(rf"\int_{{{t['T_fuel_surface']:.1f}{DEG}}}^{{2000{DEG}}} k_f\,dT = \frac{{q'}}{{4\pi}} = \frac{{{nl(qf)}}}{{4\pi}} = {qf / (4 * math.pi):.0f}\ \mathrm{{W/m}},")
    md.eq(rf"\overline{{k}}_f = \frac{{{qf / (4 * math.pi):.0f}}}{{2000 - {t['T_fuel_surface']:.1f}}} = {r['k_fuel_eff']:.2f}\ \mathrm{{W/m\,K}},")
    md.eq(rf"R_{{fuel}} = \frac{{1}}{{4\pi \times {r['k_fuel_eff']:.2f}}} = {tsci(r['R_fuel'])}\ \mathrm{{K\,m/W}}.")
    md.table(["Pass", "$q'$ guess (kW/m)", "$T_s$ (°C)", "$\\int k_f\\,dT$ (W/m)", "New $q'$ (kW/m)"],
             [[i + 1, f"{q / 1e3:.2f}", f"{Ts:.1f}", f"{qn / (4 * math.pi):.0f}", f"{qn / 1e3:.2f}"] for i, (q, Ts, qn) in enumerate(hist[:4])] +
             [["…", "", "", "", ""], [len(hist), f"{hist[-1][0] / 1e3:.2f}", f"{hist[-1][1]:.1f}", f"{hist[-1][2] / (4 * math.pi):.0f}", f"{hist[-1][2] / 1e3:.2f}"]],
             "Fixed-point iteration for the pellet-centre limit.", widths=[0.6, 1.2, 1, 1.4, 1.2])
    md.p(f"**Clad limit.** Here the clad is at about 1200 °C, where $k_c$ = {k_c_c:.1f} W/m K and $R_{{clad}}$ = "
         f"{sci(R_clad_c)} K m/W, so from Eq. ({e['lim_c']})")
    md.eq(rf"q'_{{clad}} = \frac{{1200 - 310}}{{{tsci(r['R_film'])} + {tsci(R_clad_c)}}} = {nl(r['q_clad'])}\ \mathrm{{W/m}} "
          rf"\approx {r['q_clad'] / 1e3:.0f}\ \mathrm{{kW/m}}.")
    md.p(f"**Temperatures at $q'_{{fuel}}$.** Working outwards from the coolant,")
    md.eq(rf"T_{{co}} = 310 + {nl(qf)} \times {tsci(r['R_film'])} = {t['T_clad_out']:.1f}{DEG},")
    md.eq(rf"T_{{ci}} = T_s = {t['T_clad_out']:.1f} + {nl(qf)} \times {tsci(r['R_clad'])} = {t['T_clad_in']:.1f}{DEG},")
    md.eq(rf"T_c = {t['T_fuel_surface']:.1f} + {nl(qf)} \times {tsci(r['R_fuel'])} = {t['T_centre']:.0f}{DEG}.")
    md.p("**Plant check.** The core average coolant temperature, 310 °C, is the mean of the primary temperatures of "
         "Question 2 (295 °C to 325 °C), so we can check that the rod is consistent with the plant. With the enthalpies "
         f"at 180 bar, $h$(325 °C) = {h325:.1f} kJ/kg and $h$(295 °C) = {h295:.1f} kJ/kg,")
    md.eq(rf"Q_{{th}} = \frac{{1000}}{{0.32}} = 3125\ \mathrm{{MW}}, \qquad \dot m = \frac{{3125\times10^6}}{{({h325:.1f} - {h295:.1f})\times10^3}} "
          rf"= {nl(pc['m_total'])}\ \mathrm{{kg/s}},")
    md.eq(rf"A_{{cell}} = \frac{{\sqrt3}}{{2}}(12.5)^2 - \frac{{\pi}}{{4}}(10)^2 = {pc['A_cell'] * 1e6:.1f}\ \mathrm{{mm^2}},")
    md.eq(rf"\dot m_{{rod}} = \rho V A_{{cell}} = {w['rho']:.1f} \times 5 \times {pc['A_cell'] * 1e6:.1f}\times10^{{-6}} = {pc['m_rod']:.3f}\ \mathrm{{kg/s}},")
    md.eq(rf"N_{{rods}} = \frac{{{nl(pc['m_total'])}}}{{{pc['m_rod']:.3f}}} \approx {nl(pc['N_rods'])},")
    md.eq(rf"q'_{{avg}} = \frac{{3125\times10^6}}{{{nl(pc['N_rods'])} \times 3}} = {pc['q_avg'] / 1e3:.1f}\ \mathrm{{kW/m}}, \qquad "
          rf"F_q = \frac{{{r['q_max'] / 1e3:.1f}}}{{{pc['q_avg'] / 1e3:.1f}}} = {pc['F_q_allowed']:.1f}.")

    md.h("Results", 2)
    md.p(f"The maximum permissible linear heat rate is $q'_{{max}}$ = **{r['q_max'] / 1e3:.1f} kW/m**, set by the 2000 °C "
         f"pellet-centre limit. The clad limit is {r['q_clad'] / 1e3:.0f} kW/m, about {r['q_clad'] / r['q_fuel']:.0f} times higher, "
         f"so it never comes into play. Table {md.ntab + 1} lists the temperatures at each limit and Table {md.ntab + 2} the "
         "sensitivity to the main assumptions.")
    md.p(f"At {r['q_max'] / 1e3:.1f} kW/m the clad outer surface is at {t['T_clad_out']:.0f} °C, {w['T_sat'] - t['T_clad_out']:.0f} K "
         f"below saturation, so the coolant stays single phase; the surface heat flux is {t['q_flux'] / 1e6:.2f} MW/m². Of the "
         f"{t['T_centre'] - t['T_bulk']:.0f} K between the centre and the coolant, {t['T_centre'] - t['T_fuel_surface']:.0f} K "
         f"({100 * (t['T_centre'] - t['T_fuel_surface']) / (t['T_centre'] - t['T_bulk']):.0f} %) are inside the pellet. This is "
         "why the film correlation and the coolant temperature hardly matter. A pellet–clad gap is the one thing that would "
         f"change the answer: with $h_{{gap}}$ = {npr(1e4)} W/m² K the limit drops to {g1:.0f} kW/m, and with 5000 W/m² K "
         f"to {g2:.0f} kW/m.")
    md.p(f"The plant check gives about {npr(pc['N_rods'])} rods in the core and an average linear heat rate of "
         f"{pc['q_avg'] / 1e3:.1f} kW/m, so the limit allows a total peaking factor of {pc['F_q_allowed']:.1f}. A PWR runs "
         "at about 2.5, so the rod has margin.")
    md.table(["Limit", "$q'$ (kW/m)", "$T_{co}$ (°C)", "$T_{ci}$ (°C)", "$T_s$ (°C)", "$T_c$ (°C)"], [
        ["Pellet centre 2000 °C", f"{r['q_fuel'] / 1e3:.1f}", f"{t['T_clad_out']:.0f}", f"{t['T_clad_in']:.0f}", f"{t['T_fuel_surface']:.0f}", f"{t['T_centre']:.0f}"],
        ["Clad 1200 °C", f"{r['q_clad'] / 1e3:.0f}", f"{tc['T_clad_out']:.0f}", f"{tc['T_clad_in']:.0f}", f"{tc['T_fuel_surface']:.0f}", "> 2840 (melts)"],
    ], "The two limits and the temperatures at each.", widths=[2, 1, 1, 1, 1, 1.2])
    md.table(["Case", "$q'_{max}$ (kW/m)"], [[n, f"{q / 1e3:.1f}"] for n, q in rows],
             "Sensitivity of $q'_{max}$ to the assumptions.", widths=[2.5, 1])
    md.h("Graphs", 2)
    md.fig(plots.q1_radial_profile(r, os.path.join(FIG, "Q1_radial_profile.png"), fr),
           "Radial temperature profile in the rod at $q'_{max}$.", "4.6in")
    md.p(f"Figure {md.nfig} shows the profile at $q'_{{max}}$. It is parabolic in the pellet and nearly flat across the clad "
         f"and the film. Figure {md.nfig + 1} splits the temperature drop between the layers and shows the sensitivity cases "
         f"of Table {md.ntab}.")
    md.fig(plots.q1_breakdown(r, rows, os.path.join(FIG, "Q1_breakdown.png"), fr),
           "(a) Temperature drop in each layer at $q'_{max}$; (b) sensitivity of $q'_{max}$.", "6.3in")
    md.h("Conclusion", 2)
    md.p(f"The maximum permissible linear heat rate is about {r['q_max'] / 1e3:.0f} kW/m ({r['q_max'] / 1e3:.1f} kW/m), set by "
         "the 2000 °C pellet-centre limit. The clad stays near 400 °C and is nowhere near 1200 °C, which is really a "
         "loss-of-coolant criterion; in normal operation the clad limit would be DNB, not temperature. The answer is "
         "insensitive to the film coefficient and the coolant temperature. A pellet–clad gap is the one thing that would "
         "reduce it noticeably.")
    return r


# ----------------------------------------------------------------------------- 4 Question 2
def section_q2(md, e):
    spec = dict(sg.SPEC_GIVEN, P_steam_bar=70.0)
    alts, P_max = sg.alternatives(sg.SPEC_GIVEN)
    res = sg.design(spec, sg.DESIGN)
    eb, g, dsg = res["eb"], res["geom"], res["design"]
    pa_given = sg.pinch_analysis(sg.SPEC_GIVEN)
    pa = res["pinch"]
    trade = sg.tradeoff(spec, V_list=(2.5, 3.0, 3.5, 4.0, 5.0, 6.0))
    conv = sg.convergence(spec, ns=(20, 40, 80))
    p, q, w = res["dP_shell_parts"], res["dP_tube_parts"], res["weights"]
    s = res["segments"]
    top = res["report"]["top"]
    sat70, sat100 = sg.sat_at(70.0), sg.sat_at(100.0)
    eb100 = pa_given["eb"]
    # hand calculation of the first segment with its inlet state
    dz = s["dz"][0]
    pri = sg.state_PT(180.0, s["T_pri"][0])
    liq = sg.state_Ph(eb["P_shell_in"], eb["h_fw"])
    Re_t = g["G_t"] * g["d_i"] / pri["mu"]
    Nu_t = sg.nu_dittus_boelter(Re_t, pri["Pr"], 0.3)
    h_i = Nu_t * pri["k"] / g["d_i"]
    Re_s = g["G_s"] * g["D_e"] / liq["mu"]
    Nu_s = sg.nu_dittus_boelter(Re_s, liq["Pr"], 0.4)
    h_o = Nu_s * liq["k"] / g["D_e"]
    ratio = g["d_o"] / g["d_i"]
    R_in = ratio * (1 / h_i + dsg["R_foul_in"]) + g["R_wall"] + dsg["R_foul_out"]
    T_w = (s["T_pri"][0] / R_in + h_o * s["T_sec"][0]) / (1 / R_in + h_o)
    q1 = h_o * (T_w - s["T_sec"][0])
    U1 = q1 / (s["T_pri"][0] - s["T_sec"][0])
    dQ1 = q1 * g["A_ht_per_m"] * dz
    f1 = sg.friction_factor(Re_s)
    dp_f = f1 * dz / g["D_e"] * g["G_s"] ** 2 / (2 * liq["rho"])
    dp_g = liq["rho"] * 9.81 * dz
    A_cell = g["A_cell"]
    P_sh = dsg["design_factor"] * eb["P_shell_in"] * 1e5
    L_shell = res["L_tube"] + 2 * res["t_tubesheet"]
    d_o, d_i, N = g["d_o"] * 1e3, g["d_i"] * 1e3, g["N"]
    Pd = g["P_tube_design"] / 1e6                      # MPa
    N_raw = eb["m_pri"] / (g["rho_pri_mean"] * dsg["V_tube"] * g["A_i"])
    Q_boil100 = pa_given["frac_boil"] * eb100["Q_sg"]
    h_pb100 = (eb["h_pri_out"] + Q_boil100 / eb["m_pri"]) / 1e3
    h_pb70 = (eb["h_pri_out"] + pa["frac_boil"] * eb["Q_sg"] / eb["m_pri"]) / 1e3
    Ts1 = sg.state_Ph(eb["P_shell_in"], eb["h_fw"] + dQ1 / eb["m_sec"])["T"]
    Tp1 = sg.state_Ph(180.0, eb["h_pri_out"] + dQ1 / eb["m_pri"])["T"]
    kJ = lambda x: f"{x / 1e3:.1f}"
    Tp0, Ts0 = s["T_pri"][0], s["T_sec"][0]
    dh_tot = (eb["h_sec_out"] - eb["h_fw"]) / 1e3

    md.h("Question 2: Once-through steam generator", 1)
    md.h("Calculation", 2)
    md.h("Energy balance and feasibility", 3)
    md.p(f"The thermal power of the plant and the duty of one steam generator follow from Eq. ({e['eb_q']}):")
    md.eq(rf"Q_{{th}} = \frac{{1000}}{{0.32}} = {eb['Q_th'] / 1e6:.0f}\ \mathrm{{MW}}, \qquad Q_{{SG}} = \frac{{3125}}{{4}} = {eb['Q_sg'] / 1e6:.2f}\ \mathrm{{MW}}.")
    md.p(f"On the primary side at 180 bar, $h$(325 °C) = {kJ(eb['h_pri_in'])} kJ/kg and $h$(295 °C) = {kJ(eb['h_pri_out'])} kJ/kg, "
         f"so Eq. ({e['eb_mp']}) gives")
    md.eq(rf"\dot m_p = \frac{{781.25\times10^3}}{{{kJ(eb['h_pri_in'])} - {kJ(eb['h_pri_out'])}}} = {nl(eb['m_pri'])}\ \mathrm{{kg/s}}.")
    md.p(f"As specified, with steam at 100 bar, $T_{{sat}}$ = {sat100['Tsat']:.1f} °C. The feed water then enters at "
         f"{eb100['T_fw']:.1f} °C with $h_{{fw}}$ = {kJ(eb100['h_fw'])} kJ/kg, $h_f$ = {kJ(sat100['hf'])} kJ/kg, and the steam at "
         f"320 °C has $h$ = {kJ(eb100['h_sec_out'])} kJ/kg. From Eqs. ({e['eb_ms']}) and ({e['tq_p']}),")
    md.eq(rf"\dot m_s = \frac{{781.25\times10^3}}{{{kJ(eb100['h_sec_out'])} - {kJ(eb100['h_fw'])}}} = {eb100['m_sec']:.0f}\ \mathrm{{kg/s}},")
    md.eq(rf"Q_{{boil}} = \dot m_s\,(h_f - h_{{fw}}) = {eb100['m_sec']:.0f} \times ({kJ(sat100['hf'])} - {kJ(eb100['h_fw'])}) \times 10^{{-3}} "
          rf"= {Q_boil100 / 1e6:.0f}\ \mathrm{{MW}},")
    md.eq(rf"h_p(Q_{{boil}}) = {kJ(eb['h_pri_out'])} + \frac{{{Q_boil100 / 1e6:.0f}\times10^3}}{{{nl(eb['m_pri'])}}} = {h_pb100:.1f}\ \mathrm{{kJ/kg}} "
          rf"\;\Rightarrow\; T_p = {pa_given['T_pri_boil']:.1f}{DEG},")
    md.eq(rf"\Delta T_{{boil}} = {pa_given['T_pri_boil']:.1f} - {sat100['Tsat']:.1f} = {pa_given['dT_boil']:.1f}\ \mathrm{{K}}.")
    md.cont(f"The secondary would have to boil {abs(pa_given['dT_boil']):.1f} K above the primary temperature, which is "
            f"impossible. Figure {md.nfig + 1}(a) shows the two profiles crossing. No correlation or surface area can fix "
            "this; it is fixed by the stream data alone.")
    md.p("So we keep every primary-side value, the 320 °C outlet and the 20 K subcooling, and lower the steam pressure. "
         f"For an approach of 10 K at the start of boiling the largest pressure is {P_max:.1f} bar, and we take **70 bar**. "
         f"Table {md.ntab + 1} lists the other ways out that we considered: raising the primary temperatures would bring the "
         "core outlet within 12 K of saturation at 180 bar, and keeping 100 bar with a primary outlet of 318 °C would need "
         "4.3 times the primary flow.")
    md.p(f"At 70 bar, $T_{{sat}}$ = {sat70['Tsat']:.1f} °C. The feed water enters at {eb['T_fw']:.1f} °C with $h_{{fw}}$ = "
         f"{kJ(eb['h_fw'])} kJ/kg, $h_f$ = {kJ(sat70['hf'])} kJ/kg, $h_g$ = {kJ(sat70['hg'])} kJ/kg and $h$(320 °C) = "
         f"{kJ(eb['h_sec_out'])} kJ/kg. Then")
    md.eq(rf"\dot m_s = \frac{{781.25\times10^3}}{{{kJ(eb['h_sec_out'])} - {kJ(eb['h_fw'])}}} = {eb['m_sec']:.0f}\ \mathrm{{kg/s}},")
    md.cont("and the duty splits between preheating, boiling and superheating as")
    md.eq(rf"f_{{pre}} = \frac{{h_f - h_{{fw}}}}{{h_{{s,out}} - h_{{fw}}}} = \frac{{{kJ(sat70['hf'])} - {kJ(eb['h_fw'])}}}{{{dh_tot:.1f}}} = {eb['f_pre']:.3f},")
    md.eq(rf"f_{{boil}} = \frac{{{kJ(sat70['hg'])} - {kJ(sat70['hf'])}}}{{{dh_tot:.1f}}} = {eb['f_boil']:.3f}, \qquad "
          rf"f_{{sh}} = \frac{{{kJ(eb['h_sec_out'])} - {kJ(sat70['hg'])}}}{{{dh_tot:.1f}}} = {eb['f_sh']:.3f}.")
    md.cont("The primary temperature where boiling starts and the approach there are then")
    md.eq(rf"h_p(Q_{{boil}}) = {kJ(eb['h_pri_out'])} + {pa['frac_boil']:.3f}\,({kJ(eb['h_pri_in'])} - {kJ(eb['h_pri_out'])}) = {h_pb70:.1f}\ \mathrm{{kJ/kg}} "
          rf"\;\Rightarrow\; T_p = {pa['T_pri_boil']:.1f}{DEG},")
    md.eq(rf"\Delta T_{{boil}} = {pa['T_pri_boil']:.1f} - {sat70['Tsat']:.1f} = {pa['dT_boil']:.1f}\ \mathrm{{K}}.")
    md.cont(f"At the hot end the approach is $325 - 320 = 5$ K. Figure {md.nfig + 1}(b) shows the profiles at 70 bar.")
    md.table(["Option", "Steam (bar)", "$T_{sat}$ (°C)", "Primary in / out (°C)", "Minimum approach", "$\\dot m_p$ (kg/s)"],
             [[a["name"], f"{a['P']:.0f}", f"{a['Tsat']:.1f}", f"{a['T_in']:.0f} / {a['T_out']:.0f}",
               f"{a['dT_min']:.1f} K at {100 * a['frac']:.0f} %", f"{npr(a['m_pri'])}"] for a in alts],
             "Options for the specification. The minimum approach includes the 5 K at the hot end.",
             widths=[3.8, 0.8, 0.8, 1.3, 1.5, 0.9])
    md.fig(plots.q2_tq_diagram(pa_given, pa, os.path.join(FIG, "Q2_TQ_diagram.png")),
           "Temperature–duty diagrams: (a) as specified at 100 bar, (b) at 70 bar.", "6.3in")

    md.h("Tubes and shell-side geometry", 3)
    md.p(f"We use 5/8 in tubes, $d_o$ = {d_o:.3f} mm. The design pressure is $1.1 \\times 180$ = {g['P_tube_design'] / 1e5:.0f} bar "
         f"with the shell depressurised, so Eq. ({e['tube']}) gives")
    md.eq(rf"t_{{tube}} = \frac{{{Pd:.1f} \times {d_o:.3f}}}{{2 \times 115 + 0.8 \times {Pd:.1f}}} = {g['t_tube_min'] * 1e3:.2f}\ \mathrm{{mm}} "
          rf"\;\rightarrow\; \text{{BWG {g['gauge']}}},\ t = {g['t_tube'] * 1e3:.3f}\ \mathrm{{mm}},")
    md.eq(rf"d_i = {d_o:.3f} - 2 \times {g['t_tube'] * 1e3:.3f} = {d_i:.2f}\ \mathrm{{mm}}, \qquad A_i = \frac{{\pi}}{{4}}\,d_i^2 = {g['A_i'] * 1e6:.1f}\ \mathrm{{mm^2}}.")
    md.p(f"We take a tube-side velocity of 3.5 m/s (Table {md.ntab + 5} shows the trade-off; at 5 m/s the tubes would be "
         f"24 m long). With $\\rho_p$ = {g['rho_pri_mean']:.0f} kg/m³ at the mean primary temperature of 310 °C, "
         f"Eq. ({e['N']}) gives")
    md.eq(rf"N = \frac{{{nl(eb['m_pri'])}}}{{{g['rho_pri_mean']:.0f} \times 3.5 \times {g['A_i'] * 1e6:.1f}\times10^{{-6}}}} = {nl(N_raw, 1)} "
          rf"\;\rightarrow\; {nl(N)}\ \text{{tubes}},")
    md.eq(rf"G_t = \frac{{{nl(eb['m_pri'])}}}{{{nl(N)} \times {g['A_i'] * 1e6:.1f}\times10^{{-6}}}} = {g['G_t']:.0f}\ \mathrm{{kg/m^2\,s}}.")
    md.p(f"On a triangular pitch of $1.3\\,d_o$ = {g['pitch'] * 1e3:.2f} mm, Eqs. ({e['afree']})–({e['de2']}) give")
    md.eq(rf"A_{{free}} = \frac{{\sqrt3}}{{2}}({g['pitch'] * 1e3:.2f})^2 - \frac{{\pi}}{{4}}({d_o:.3f})^2 = {g['A_free'] * 1e6:.1f}\ \mathrm{{mm^2}},")
    md.eq(rf"G_s = \frac{{{eb['m_sec']:.0f}}}{{{nl(N)} \times {g['A_free'] * 1e6:.1f}\times10^{{-6}}}} = {g['G_s']:.0f}\ \mathrm{{kg/m^2\,s}},")
    md.eq(rf"D_e = {d_o:.3f}\left[\frac{{2\sqrt3}}{{\pi}}(1.3)^2 - 1\right] = {g['D_e'] * 1e3:.2f}\ \mathrm{{mm}},")
    md.cont("and the wall term of the resistance is")
    md.eq(rf"R_{{wall}} = \frac{{d_o\ln(d_o/d_i)}}{{2k_w}} = \frac{{{d_o:.3f}\times10^{{-3}} \times \ln {ratio:.4f}}}{{2 \times 17.5}} "
          rf"= {tsci(g['R_wall'])}\ \mathrm{{m^2\,K/W}}.")

    md.h("One segment worked by hand", 3)
    md.p(f"To show what happens in each segment, here is the first one, from $z$ = 0 to {dz:.3f} m, worked with the inlet "
         f"state of the segment. The code uses the mid-segment state, so its numbers in Table {md.ntab + 4} differ a little.")
    md.p(f"Tube side: primary at {Tp0:.1f} °C and 180 bar, with $\\rho$ = {pri['rho']:.1f} kg/m³, $\\mu$ = "
         f"{pri['mu'] * 1e6:.1f} µPa s, $k$ = {pri['k']:.3f} W/m K and Pr = {pri['Pr']:.3f}. From Eq. ({e['db']}),")
    md.eq(rf"\mathrm{{Re}} = \frac{{G_t d_i}}{{\mu}} = \frac{{{g['G_t']:.0f} \times {d_i:.2f}\times10^{{-3}}}}{{{pri['mu'] * 1e6:.1f}\times10^{{-6}}}} "
          rf"= {tsci(Re_t)},")
    md.eq(rf"\mathrm{{Nu}} = 0.023\,({tsci(Re_t)})^{{0.8}}\,({pri['Pr']:.3f})^{{0.3}} = {Nu_t:.0f},")
    md.eq(rf"h_i = \frac{{\mathrm{{Nu}}\,k}}{{d_i}} = \frac{{{Nu_t:.0f} \times {pri['k']:.3f}}}{{{d_i:.2f}\times10^{{-3}}}} = {nl(h_i)}\ \mathrm{{W/m^2\,K}}.")
    md.p(f"Shell side: liquid at {Ts0:.1f} °C and {eb['P_shell_in']:.1f} bar, with $\\rho$ = {liq['rho']:.1f} kg/m³, "
         f"$\\mu$ = {liq['mu'] * 1e6:.1f} µPa s, $k$ = {liq['k']:.3f} W/m K and Pr = {liq['Pr']:.3f}:")
    md.eq(rf"\mathrm{{Re}} = \frac{{G_s D_e}}{{\mu}} = \frac{{{g['G_s']:.0f} \times {g['D_e'] * 1e3:.2f}\times10^{{-3}}}}{{{liq['mu'] * 1e6:.1f}\times10^{{-6}}}} "
          rf"= {tsci(Re_s)},")
    md.eq(rf"\mathrm{{Nu}} = 0.023\,({tsci(Re_s)})^{{0.8}}\,({liq['Pr']:.3f})^{{0.4}} = {Nu_s:.1f},")
    md.eq(rf"h_o = \frac{{{Nu_s:.1f} \times {liq['k']:.3f}}}{{{g['D_e'] * 1e3:.2f}\times10^{{-3}}}} = {nl(h_o)}\ \mathrm{{W/m^2\,K}}.")
    md.p(f"Wall temperature and heat flux, from Eqs. ({e['rin']}) and ({e['q']}) with the fouling on both sides:")
    md.eq(rf"R_{{in}} = {ratio:.4f}\left(\frac{{1}}{{{nl(h_i)}}} + 2.5\times10^{{-5}}\right) + {tsci(g['R_wall'])} + 2.5\times10^{{-5}} "
          rf"= {tsci(R_in)}\ \mathrm{{m^2\,K/W}},")
    md.eq(rf"\frac{{{Tp0:.1f} - T_w}}{{{tsci(R_in)}}} = {nl(h_o)}\,(T_w - {Ts0:.1f}) \;\Rightarrow\; T_w = {T_w:.1f}{DEG},")
    md.cont(f"which is below $T_{{sat}}$ = {s['Tsat'][0]:.1f} °C, so there is no nucleate boiling yet and the Chen "
            "correlation reduces to its convective part. Then")
    md.eq(rf"q'' = {nl(h_o)} \times ({T_w:.1f} - {Ts0:.1f}) = {q1 / 1e3:.1f}\ \mathrm{{kW/m^2}},")
    md.eq(rf"U = \frac{{q''}}{{T_p - T_s}} = \frac{{{q1 / 1e3:.1f}\times10^3}}{{{Tp0:.1f} - {Ts0:.1f}}} = {U1:.0f}\ \mathrm{{W/m^2\,K}},")
    md.eq(rf"\Delta Q = {q1 / 1e3:.1f}\times10^3 \times {nl(N)} \times \pi \times {d_o:.3f}\times10^{{-3}} \times {dz:.3f} = {dQ1 / 1e6:.2f}\ \mathrm{{MW}},")
    md.eq(rf"\Delta h_s = \frac{{{dQ1 / 1e6:.2f}\times10^6}}{{{eb['m_sec']:.0f}}} = {dQ1 / eb['m_sec'] / 1e3:.1f}\ \mathrm{{kJ/kg}} "
          rf"\;\Rightarrow\; T_s = {Ts1:.1f}{DEG},")
    md.eq(rf"\Delta h_p = \frac{{{dQ1 / 1e6:.2f}\times10^6}}{{{nl(eb['m_pri'])}}} = {dQ1 / eb['m_pri'] / 1e3:.2f}\ \mathrm{{kJ/kg}} "
          rf"\;\Rightarrow\; T_p = {Tp1:.1f}{DEG}.")
    md.p(f"Pressure drop of the segment on the shell side, from Eqs. ({e['f']}) and ({e['dp']}):")
    md.eq(rf"f = (0.79 \ln {tsci(Re_s)} - 1.64)^{{-2}} = {f1:.4f},")
    md.eq(rf"\Delta p_f = {f1:.4f} \times \frac{{{dz:.3f}}}{{{g['D_e'] * 1e3:.2f}\times10^{{-3}}}} \times \frac{{{g['G_s']:.0f}^2}}{{2 \times {liq['rho']:.0f}}} "
          rf"= {dp_f:.0f}\ \mathrm{{Pa}},")
    md.eq(rf"\Delta p_g = {liq['rho']:.0f} \times 9.81 \times {dz:.3f} = {dp_g:.0f}\ \mathrm{{Pa}}.")
    md.cont("In the boiling region the homogeneous density takes the place of the liquid density, and the support-plate "
            "loss is added once per metre.")

    md.h("Shell, heads, tubesheets and weight", 3)
    md.p(f"Each tube occupies a cell of $(\\sqrt3/2)\\,p^2$ = {A_cell * 1e6:.1f} mm². With 93 % of the bundle circle used,")
    md.eq(rf"D_{{otl}} = \sqrt{{\frac{{4 \times {nl(N)} \times {A_cell * 1e6:.1f}\times10^{{-6}}}}{{\pi \times 0.93}}}} = {g['D_otl']:.2f}\ \mathrm{{m}},")
    md.eq(rf"D_s = {g['D_otl']:.2f} + 0.05 = {g['D_shell']:.2f}\ \mathrm{{m}}.")
    md.p(f"The shell design pressure is $1.1 \\times {eb['P_shell_in']:.1f}$ = {P_sh / 1e5:.1f} bar with $R$ = {g['D_shell'] / 2:.2f} m. "
         f"The heads and tubesheets see the primary design pressure of {g['P_tube_design'] / 1e5:.0f} bar. "
         f"Eqs. ({e['shell']})–({e['ts']}) give")
    md.eq(rf"t_{{shell}} = \frac{{{P_sh / 1e6:.2f} \times {g['D_shell'] / 2 * 1e3:.0f}}}{{115 - 0.6 \times {P_sh / 1e6:.2f}}} = {res['t_shell'] * 1e3:.0f}\ \mathrm{{mm}},")
    md.eq(rf"t_{{head}} = \frac{{{Pd:.1f} \times {g['D_shell'] / 2 * 1e3:.0f}}}{{2 \times 115 - 0.2 \times {Pd:.1f}}} = {res['t_head'] * 1e3:.0f}\ \mathrm{{mm}},")
    md.eq(rf"T_{{ts}} = \frac{{{g['D_shell']:.2f}}}{{3}}\sqrt{{\frac{{{Pd:.1f}}}{{115}}}} = {res['t_tubesheet'] * 1e3:.0f}\ \mathrm{{mm}}.")
    md.p(f"The march gives a tube length of $L$ = {res['L_tube']:.2f} m (40 segments of {dz:.3f} m); a fine march with 0.1 m "
         f"steps gives {res['L_fine']:.2f} m. There are {res['n_supports']} support plates, and the overall height with "
         f"tubesheets and heads is {res['H_overall']:.1f} m. The dry weight, with $\\rho$ = 8000 kg/m³ for SS316, is")
    md.eq(rf"m_{{tubes}} = {nl(N)} \times \frac{{\pi}}{{4}}\left({d_o:.3f}^2 - {d_i:.2f}^2\right)\times10^{{-6}} \times {res['L_tube']:.2f} \times 8000 "
          rf"= {w['tubes'] / 1e3:.0f}\ \mathrm{{t}},")
    md.eq(rf"m_{{shell}} = \pi\,({g['D_shell']:.2f} + {res['t_shell']:.3f}) \times {res['t_shell']:.3f} \times {L_shell:.2f} \times 8000 = {w['shell'] / 1e3:.0f}\ \mathrm{{t}},")
    md.eq(rf"m_{{ts}} = 2 \times \frac{{\pi}}{{4}}\left({g['D_shell']:.2f}^2 - {nl(N)} \times {d_o:.3f}^2\times10^{{-6}}\right) \times {res['t_tubesheet']:.3f} \times 8000 "
          rf"= {w['tubesheets'] / 1e3:.0f}\ \mathrm{{t}},")
    md.cont(f"and the heads add {w['heads'] / 1e3:.0f} t and the support plates {w['supports'] / 1e3:.0f} t. The sum is "
            f"{w['subtotal'] / 1e3:.0f} t, and with 10 % for nozzles, skirt and internals the unladen weight is "
            f"**{res['weight'] / 1e3:.0f} t**.")
    md.p("The pressure drops over the full height are the sums over the segments. On the tube side the friction and "
         f"the end losses add up to {res['dP_tube'] / 1e5:.2f} bar; the downflow also gains {-q['gravity'] / 1e5:.2f} bar of "
         "static head, which we report separately. On the shell side the four terms of Eq. "
         f"({e['dp']}) add up to {res['dP_shell'] / 1e5:.2f} bar:")
    md.eq(rf"\Delta p_{{tube}} = \Delta p_f + \Delta p_{{ends}} = {q['friction'] / 1e5:.2f} + {q['ends'] / 1e5:.2f} = {res['dP_tube'] / 1e5:.2f}\ \mathrm{{bar}},")
    md.eq(rf"\Delta p_{{shell}} = {p['friction'] / 1e5:.3f} + {p['gravity'] / 1e5:.3f} + {p['acceleration'] / 1e5:.3f} + {p['supports'] / 1e5:.3f} "
          rf"= {res['dP_shell'] / 1e5:.2f}\ \mathrm{{bar}}.")
    md.p(f"Table {md.ntab + 1} collects the assumptions behind these numbers.")
    md.table(["Item", "Value"], [
        ["Tube", f"OD {d_o:.3f} mm, BWG {g['gauge']} wall {g['t_tube'] * 1e3:.3f} mm, SS316"],
        ["Material data", f"$S$ = {dsg['S_allow'] / 1e6:.0f} MPa (ASME II-D, about 650 °F), $E$ = 1, $k_w$ = {dsg['k_wall']} W/m K, $\\rho$ = {dsg['rho_metal']:.0f} kg/m³"],
        ["Layout", f"triangular, $p/d_o$ = {dsg['pitch_ratio']}, tube-side velocity {dsg['V_tube']} m/s"],
        ["Fouling", f"{sci(dsg['R_foul_in'], 1)} m² K/W on each side"],
        ["Dryout", f"at $x$ = {dsg['x_dryout']} (look-up-table range for 7 MPa, 200 kg/m² s, 0.1 MW/m²)"],
        ["Support plates", f"every {dsg['support_spacing']} m, $K$ = {dsg['K_support']}; tube ends $K$ = {dsg['K_tube_ends']}"],
        ["Shell", f"ID = bundle diameter + {dsg['shell_clearance'] * 1e3:.0f} mm, bundle area use 0.93; hemispherical heads"],
    ], "Assumptions used in the design.", widths=[1.2, 4])

    md.h("Results", 2)
    md.p(f"One steam generator needs **{npr(res['N_tubes'])} tubes** of {d_o:.3f} mm outer diameter, "
         f"**{res['d_i'] * 1e3:.2f} mm** inner diameter and **{res['t_tube'] * 1e3:.3f} mm** wall (BWG {res['gauge']}), with a "
         f"heated height of **{res['L_tube']:.1f} m**. The shell has an inner diameter of **{res['D_shell']:.2f} m** and a wall "
         f"of **{res['t_shell'] * 1e3:.0f} mm**; the heads are {res['t_head'] * 1e3:.0f} mm and the tubesheets "
         f"{res['t_tubesheet'] * 1e3:.0f} mm thick. The pressure drop is **{res['dP_shell'] / 1e5:.2f} bar** on the shell side "
         f"and **{res['dP_tube'] / 1e5:.2f} bar** on the tube side (friction and form losses; the tube side also gains "
         f"{-q['gravity'] / 1e5:.2f} bar of static head on the way down). The unladen weight is **{res['weight'] / 1e3:.0f} t**. "
         f"Table {md.ntab + 1} lists these against the five items asked for.")
    md.p(f"The heat-transfer surface is {npr(res['A_total'])} m², the mean $U$ is {res['U_mean']:.0f} W/m² K and the mean heat "
         f"flux {res['q_mean'] / 1e3:.0f} kW/m². The primary comes out at {res['report']['T_pri_top']:.2f} °C at the top and the "
         f"steam at {res['report']['T_sec_top']:.2f} °C and {res['report']['P_sec_out']:.2f} bar, as required, which closes the "
         f"energy balance. Table {md.ntab + 2} gives the four regions and Table {md.ntab + 3} the 40 segments.")
    md.table(["Required quantity", "Result"], [
        ["(i) Number of tubes", f"{npr(res['N_tubes'])}"],
        ["(ii) Tube inner diameter / wall / height", f"{res['d_i'] * 1e3:.2f} mm / {res['t_tube'] * 1e3:.3f} mm / {res['L_tube']:.1f} m"],
        ["(iii) Shell inner diameter / wall", f"{res['D_shell']:.2f} m / {res['t_shell'] * 1e3:.0f} mm"],
        ["(iv) Pressure drop, shell side / tube side", f"{res['dP_shell'] / 1e5:.2f} bar / {res['dP_tube'] / 1e5:.2f} bar"],
        ["(v) Unladen weight", f"{res['weight'] / 1e3:.0f} t"],
    ], "Design results for one steam generator (70 bar steam).", widths=[2, 2])
    bounds = {}
    for z0, z1, reg in zip(s["z0"], s["z1"], s["regime"]):
        bounds.setdefault(reg, [z0, z1])[1] = z1
    md.table(["Region", "From (m)", "To (m)", "Share of height", "Share of duty"],
             [[reg, f"{b[0]:.2f}", f"{b[1]:.2f}", f"{100 * (b[1] - b[0]) / res['L_tube']:.0f} %",
               f"{100 * sum(dq for dq, rg in zip(s['dQ'], s['regime']) if rg == reg) / eb['Q_sg']:.0f} %"] for reg, b in bounds.items()],
             "Shell-side regions, from the bottom.", widths=[1.6, 1, 1, 1.3, 1.3])
    code = {"subcooled": "SC", "nucleate boiling": "NB", "post-dryout": "PD", "superheated": "SH"}
    md.table(["#", "z₀ (m)", "Tp (°C)", "Ts (°C)", "Tw (°C)", "x", "hi", "ho", "U", "q″ (kW/m²)", "Ps (bar)", "ΔQ (MW)", "Regime"],
             [[i + 1, f"{s['z0'][i]:.2f}", f"{s['T_pri'][i]:.1f}", f"{s['T_sec'][i]:.1f}", f"{s['T_wall_out'][i]:.1f}", f"{s['x'][i]:.3f}",
               f"{s['h_i'][i]:.0f}", f"{s['h_o'][i]:.0f}", f"{s['U'][i]:.0f}", f"{s['q'][i] / 1e3:.1f}", f"{s['P_sec'][i]:.2f}",
               f"{s['dQ'][i] / 1e6:.2f}", code[s["regime"][i]]] for i in range(len(s["z"]))] +
             [["top", f"{res['L_tube']:.2f}", f"{top['T_pri']:.1f}", f"{top['T_sec']:.1f}", f"{top['T_wall_out']:.1f}", f"{top['x']:.3f}",
               f"{top['h_i']:.0f}", f"{top['h_o']:.0f}", f"{top['U']:.0f}", f"{top['q'] / 1e3:.1f}", f"{top['P_sec']:.2f}", "", code[top["regime"]]]],
             "Segment-by-segment results (values at the inlet of each segment). Tp, Ts, Tw: primary, secondary and "
             "outer-wall temperatures; hi, ho, U in W/m² K. SC subcooled, NB nucleate boiling, PD post-dryout, SH superheated.",
             widths=[0.6, 0.9, 0.9, 0.9, 0.9, 0.9, 1.0, 1.0, 0.9, 1.1, 0.9, 0.9, 0.9])
    md.h("Graphs", 2)
    md.fig(plots.q2_axial_profiles(res, os.path.join(FIG, "Q2_axial_profiles.png")),
           "Axial profiles: (a) temperatures, (b) steam quality, (c) heat-transfer coefficients, (d) local heat flux.", "6.3in")
    md.p(f"Figure {md.nfig} shows the axial profiles. Boiling does "
         f"{100 * sum(dq for dq, rg in zip(s['dQ'], s['regime']) if rg == 'nucleate boiling') / eb['Q_sg']:.0f} % of the duty on "
         "about half the height with $U$ around 4 kW/m² K; the wall, the fouling and the boiling film each take roughly a "
         "third of the resistance. At dryout the shell-side coefficient drops by a factor of ten and the wall temperature "
         "jumps by about 20 K. The superheater needs a quarter of the height for 7 % of the duty, because the steam-side "
         "coefficient is only about 1.3 kW/m² K and the temperature difference closes to 5 K at the top.")
    md.fig(plots.q2_pressure_and_duty(res, os.path.join(FIG, "Q2_pressure_duty.png")),
           "(a) Shell-side pressure and (b) cumulative duty along the height.", "6.3in")
    md.p(f"Figure {md.nfig} shows the shell-side pressure and the cumulative duty. Most of the shell-side pressure drop is "
         "the static head of the two-phase column; the friction term is small because the shell-side mass flux is low.")
    md.table(["$V_t$ (m/s)", "$N$", "$L$ (m)", "$D_s$ (m)", "$\\Delta p$ tube (bar)", "$\\Delta p$ shell (bar)", "Weight (t)"],
             [[f"{t['V']:.1f}", npr(t["N"]), f"{t['L']:.1f}", f"{t['D_shell']:.2f}", f"{t['dP_tube'] / 1e5:.2f}", f"{t['dP_shell'] / 1e5:.2f}", f"{t['weight'] / 1e3:.0f}"] for t in trade],
             "Effect of the tube-side velocity (everything else unchanged).", widths=[1, 1, 1, 1, 1.3, 1.3, 1])
    md.table(["Segments", "$L$ (m)", "Fine march (m)", "Primary at the top (°C)", "Steam outlet (°C)"],
             [[c["n"], f"{c['L']:.2f}", f"{c['L_fine']:.2f}", f"{c['T_pri_top']:.2f}", f"{c['T_sec_top']:.2f}"] for c in conv],
             "Effect of the number of segments.")
    md.fig(plots.q2_design_space(trade, conv, dsg["V_tube"], os.path.join(FIG, "Q2_design_space.png")),
           "(a) Height and (b) weight against the tube-side velocity; (c) height against the number of segments.", "6.3in")
    md.p(f"Tables {md.ntab - 1} and {md.ntab} and Figure {md.nfig} show the two checks we made on the design choices. The "
         "surface needed hardly changes with the number of tubes, because the wall, the fouling and the boiling film set the "
         "resistance; fewer tubes just make a longer bundle with a larger tube-side pressure drop. 3.5 m/s keeps the tubes "
         "under 20 m and the tube-side pressure drop near 1 bar. With 20, 40 and 80 segments the height agrees within "
         "0.5 %, so 40 segments are enough.")
    md.h("Conclusion", 2)
    md.p("The specification cannot be met as written: 100 bar steam boils at 311 °C and the primary leaves at 295 °C. We "
         f"designed for 70 bar steam and kept every other value. One steam generator then needs {npr(res['N_tubes'])} tubes "
         f"of {d_o:.3f} mm OD ({res['d_i'] * 1e3:.2f} mm ID, {res['t_tube'] * 1e3:.3f} mm wall) and {res['L_tube']:.1f} m height "
         f"in a {res['D_shell']:.2f} m shell with a {res['t_shell'] * 1e3:.0f} mm wall. The pressure drops are "
         f"{res['dP_shell'] / 1e5:.2f} bar on the shell side and {res['dP_tube'] / 1e5:.2f} bar in the tubes, and the dry weight "
         f"is about {res['weight'] / 1e3:.0f} t. Nucleate boiling does most of the work; the post-dryout and superheat regions "
         "take over a third of the height for a sixth of the duty, because the steam-side coefficients are low and the "
         "hot-end approach is only 5 K.")
    return res


# ----------------------------------------------------------------------------- 5, 6, 7
def section_summary(md, r, res):
    md.h("Summary and limitations", 1)
    md.p(f"Table {md.ntab + 1} collects the answers to both questions.")
    md.table(["Q", "Quantity", "Result"], [
        ["1", "Maximum permissible linear heat rate", f"{r['q_max'] / 1e3:.1f} kW/m (pellet-centre limit; clad limit {r['q_clad'] / 1e3:.0f} kW/m)"],
        ["2", "(i) Number of tubes", f"{npr(res['N_tubes'])}"],
        ["2", "(ii) Tube inner diameter / wall / height", f"{res['d_i'] * 1e3:.2f} mm / {res['t_tube'] * 1e3:.3f} mm / {res['L_tube']:.1f} m"],
        ["2", "(iii) Shell inner diameter / wall", f"{res['D_shell']:.2f} m / {res['t_shell'] * 1e3:.0f} mm"],
        ["2", "(iv) Pressure drop, shell / tube", f"{res['dP_shell'] / 1e5:.2f} bar / {res['dP_tube'] / 1e5:.2f} bar"],
        ["2", "(v) Unladen weight", f"{res['weight'] / 1e3:.0f} t"],
    ], "Summary of results.", widths=[0.4, 2.2, 3])
    md.p("The main limitations of the two calculations are these.")
    md.bullets([
        "Question 1 assumes no pellet–clad gap, because none is given. A gap conductance of 5000–10 000 W/m² K would "
        "lower the limit to 41–47 kW/m. In service the clad limit would be DNB rather than 1200 °C.",
        "Question 2 uses 70 bar steam instead of 100 bar, because 100 bar cannot be reached with a primary outlet of 295 °C.",
        "The dryout quality (0.9) is taken from look-up-table values, not calculated; the post-dryout length is about 13 % "
        "of the height and scales with $(1 - x_{do})$.",
        "The shell side is treated as uniform axial flow through the bundle. A real unit has a downcomer and feed nozzles, "
        "which change the subcooled region but not the surface.",
        "SS316 with 115 MPa allowable stress gives a thick tube wall; Inconel 690 tubes would be thinner and shorten the "
        "bundle by roughly 20 %. The weight is dry metal from thin-wall formulas plus 10 % for nozzles and internals.",
    ])


def section_references(md):
    md.h("References", 1)
    md.refs([
        "N. E. Todreas and M. S. Kazimi, *Nuclear Systems I: Thermal Hydraulic Fundamentals*, 2nd ed., CRC Press, 2012.",
        "J. Weisman, \"Heat transfer to water flowing parallel to tube bundles,\" *Nuclear Science and Engineering*, vol. 6, "
        "pp. 78–79, 1959.",
        "IAPWS, *Revised Release on the IAPWS Industrial Formulation 1997 for the Thermodynamic Properties of Water and "
        "Steam (IAPWS-IF97)*, 2007.",
        "D. L. Hagrman and G. A. Reymann (eds.), *MATPRO-Version 11: A Handbook of Materials Properties for Use in the "
        "Analysis of Light Water Reactor Fuel Rod Behavior*, NUREG/CR-0497, 1979.",
        "J. K. Fink, \"Thermophysical properties of uranium dioxide,\" *Journal of Nuclear Materials*, vol. 279, pp. 1–18, 2000.",
        "ASME, *Boiler and Pressure Vessel Code*, Section VIII Division 1 and Section II Part D, ASME, New York, 2021.",
        "F. W. Dittus and L. M. K. Boelter, \"Heat transfer in automobile radiators of the tubular type,\" *University of "
        "California Publications in Engineering*, vol. 2, pp. 443–461, 1930.",
        "J. C. Chen, \"Correlation for boiling heat transfer to saturated fluids in convective flow,\" *Industrial & "
        "Engineering Chemistry Process Design and Development*, vol. 5, no. 3, pp. 322–329, 1966.",
        "R. S. Dougall and W. M. Rohsenow, \"Film boiling on the inside of vertical tubes with upward flow of the fluid at "
        "low qualities,\" MIT Report 9079-26, 1963.",
        "G. K. Filonenko, \"Friction factor for turbulent pipe flow,\" *Teploenergetika*, vol. 1, no. 4, pp. 40–44, 1954.",
    ])


def section_code(md):
    md.h("Code", 1)
    md.p("The two Colab notebooks are printed below without outputs; the plotting cells are left out. The first cell "
         "installs the iapws package, which provides the IAPWS-IF97 steam tables.")
    md.h("Notebook Q1_fuel_rod_heat_rate.ipynb", 2)
    notebook_markdown(os.path.join(NB_DIR, "Q1_fuel_rod_heat_rate.ipynb"), md)
    md.h("Notebook Q2_once_through_steam_generator.ipynb", 2)
    notebook_markdown(os.path.join(NB_DIR, "Q2_once_through_steam_generator.ipynb"), md)


# ----------------------------------------------------------------------------- build
def build(md, out):
    md_path = os.path.join(SUB, "report.md")
    with open(md_path, "w", encoding="utf-8") as f:
        f.write(md.text())
    ref = reference_docx(os.path.join(SUB, "reference.docx"))
    subprocess.run(["pandoc", md_path, "-o", out, "--reference-doc", ref, "--resource-path", SUB, "--syntax-highlighting=none",
                    "--number-sections",
                    "-f", "markdown+tex_math_dollars+pipe_tables+implicit_figures+link_attributes+fenced_code_blocks+fenced_divs"],
                   check=True)
    os.remove(ref)
    finish_docx(out)
    return out


def main(out=None):
    out = out or os.path.join(SUB, "report.docx")
    os.makedirs(FIG, exist_ok=True)
    md = MD()
    md.lines += ["---", 'title: "Reactor Thermal Hydraulics"',
                 'subtitle: "Fuel-rod heat-rate limit and once-through steam generator design"',
                 'date: "October 2026"', "---", ""]
    section_problem(md)
    e = section_method(md)
    r = section_q1(md, e)
    res = section_q2(md, e)
    section_summary(md, r, res)
    section_references(md)
    section_code(md)
    build(md, out)
    print(out)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else None)
